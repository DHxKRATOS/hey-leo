import logging
import asyncio
from typing import List, Optional, Dict, Any
import uuid
from io import BytesIO

import PyPDF2
from docx import Document
from app.models.vector import Chunk, SourceType
from app.core.exceptions import DocumentProcessingException
from app.core.config import settings

logger = logging.getLogger(__name__)

class DocumentService:
    """Service for processing documents and converting them to chunks"""
    
    def __init__(self):
        self.max_chunk_size = settings.MAX_CHUNK_SIZE
        self.chunk_overlap = settings.CHUNK_OVERLAP
        self.page_as_chunk = settings.PAGE_AS_CHUNK
    
    def _split_text_into_chunks(self, text: str, metadata: Dict[str, Any] = None, document_id: str = None, chunk_offset: int = 0) -> List[Chunk]:
        """Split text into chunks based on configuration"""
        if not text.strip():
            return []
        
        chunks = []
        
        if self.page_as_chunk:
            # Treat entire text as one chunk (for page-based chunking)
            chunks.append(Chunk(
                text=text.strip(),
                metadata={**(metadata or {}), "chunk_index": chunk_offset},
                document_id=document_id,
                chunk_index=chunk_offset
            ))
            return chunks
        else:
            # Split into smaller chunks
            words = text.split()
            current_chunk = []
            current_size = 0
            chunk_index = 0
            
            for word in words:
                word_size = len(word) + 1  # +1 for space
                
                if current_size + word_size > self.max_chunk_size and current_chunk:
                    # Create chunk from current words
                    chunk_text = " ".join(current_chunk)
                    chunks.append(Chunk(
                        text=chunk_text,
                        metadata={**(metadata or {}), "chunk_index": chunk_offset + chunk_index},
                        document_id=document_id,
                        chunk_index=chunk_offset + chunk_index
                    ))
                    
                    # Start new chunk with overlap
                    overlap_words = current_chunk[-self.chunk_overlap:] if len(current_chunk) > self.chunk_overlap else current_chunk
                    current_chunk = overlap_words + [word]
                    current_size = sum(len(w) + 1 for w in current_chunk)
                    chunk_index += 1
                else:
                    current_chunk.append(word)
                    current_size += word_size
            
            # Add final chunk if not empty
            if current_chunk:
                chunk_text = " ".join(current_chunk)
                chunks.append(Chunk(
                    text=chunk_text,
                    metadata={**(metadata or {}), "chunk_index": chunk_offset + chunk_index},
                    document_id=document_id,
                    chunk_index=chunk_offset + chunk_index
                ))
        
        return chunks
    
    async def process_pdf(self, file_content: bytes, filename: str, document_id: str = None) -> List[Chunk]:
        """Process PDF file and extract text chunks"""
        try:
            # Generate single document ID for all chunks
            if document_id is None:
                document_id = str(uuid.uuid4())
            
            pdf_file = BytesIO(file_content)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            
            chunks = []
            chunk_counter = 0
            
            for page_num, page in enumerate(pdf_reader.pages):
                try:
                    text = page.extract_text()
                    if text.strip():
                        page_metadata = {
                            "source_type": SourceType.DOCUMENT,
                            "filename": filename,
                            "page_number": page_num + 1,
                            "total_pages": len(pdf_reader.pages),
                            "document_id": document_id
                        }
                        
                        if self.page_as_chunk:
                            # Each page is a separate chunk
                            chunks.append(Chunk(
                                text=text.strip(),
                                metadata=page_metadata,
                                document_id=document_id,
                                chunk_index=chunk_counter
                            ))
                            chunk_counter += 1
                        else:
                            # Split page text into smaller chunks
                            page_chunks = self._split_text_into_chunks(text, page_metadata, document_id, chunk_counter)
                            chunk_counter += len(page_chunks)
                            chunks.extend(page_chunks)
                            
                except Exception as e:
                    logger.warning(f"Error processing page {page_num + 1} of {filename}: {str(e)}")
                    continue
            
            logger.info(f"Processed PDF {filename}: {len(chunks)} chunks from {len(pdf_reader.pages)} pages")
            return chunks
            
        except Exception as e:
            logger.error(f"Error processing PDF {filename}: {str(e)}")
            raise DocumentProcessingException(f"Failed to process PDF: {str(e)}")
    
    async def process_docx(self, file_content: bytes, filename: str, document_id: str = None) -> List[Chunk]:
        """Process DOCX file and extract text chunks"""
        try:
            # Generate single document ID for all chunks
            if document_id is None:
                document_id = str(uuid.uuid4())
            
            docx_file = BytesIO(file_content)
            doc = Document(docx_file)
            
            # Extract all text from paragraphs
            full_text = ""
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    full_text += paragraph.text + "\n"
            
            if not full_text.strip():
                logger.warning(f"No text found in DOCX file {filename}")
                return []
            
            metadata = {
                "source_type": SourceType.DOCUMENT,
                "filename": filename,
                "total_paragraphs": len(doc.paragraphs),
                "document_id": document_id
            }
            
            chunks = self._split_text_into_chunks(full_text, metadata, document_id, 0)
            logger.info(f"Processed DOCX {filename}: {len(chunks)} chunks from {len(doc.paragraphs)} paragraphs")
            return chunks
            
        except Exception as e:
            logger.error(f"Error processing DOCX {filename}: {str(e)}")
            raise DocumentProcessingException(f"Failed to process DOCX: {str(e)}")
    
    async def process_txt(self, file_content: bytes, filename: str, document_id: str = None) -> List[Chunk]:
        """Process TXT file and extract text chunks"""
        try:
            # Generate single document ID for all chunks
            if document_id is None:
                document_id = str(uuid.uuid4())
            
            # Try different encodings
            text = None
            encodings = ['utf-8', 'latin-1', 'cp1252']
            
            for encoding in encodings:
                try:
                    text = file_content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            
            if text is None:
                raise DocumentProcessingException(f"Could not decode text file {filename} with any supported encoding")
            
            if not text.strip():
                logger.warning(f"No text found in TXT file {filename}")
                return []
            
            metadata = {
                "source_type": SourceType.DOCUMENT,
                "filename": filename,
                "file_size": len(file_content),
                "document_id": document_id
            }
            
            chunks = self._split_text_into_chunks(text, metadata, document_id, 0)
            logger.info(f"Processed TXT {filename}: {len(chunks)} chunks")
            return chunks
            
        except Exception as e:
            logger.error(f"Error processing TXT {filename}: {str(e)}")
            raise DocumentProcessingException(f"Failed to process TXT: {str(e)}")
    
    async def process_file(self, file_content: bytes, filename: str, document_id: str = None) -> List[Chunk]:
        """Process file based on extension"""
        # Generate single document ID for all chunks if not provided
        if document_id is None:
            document_id = str(uuid.uuid4())
        
        file_extension = filename.lower().split('.')[-1]
        
        if file_extension == 'pdf':
            return await self.process_pdf(file_content, filename, document_id)
        elif file_extension == 'docx':
            return await self.process_docx(file_content, filename, document_id)
        elif file_extension == 'txt':
            return await self.process_txt(file_content, filename, document_id)
        else:
            raise DocumentProcessingException(f"Unsupported file type: {file_extension}")

# Singleton instance
_document_service_instance = None

async def get_document_service() -> DocumentService:
    """Get or create document service instance"""
    global _document_service_instance
    if _document_service_instance is None:
        _document_service_instance = DocumentService()
    return _document_service_instance
